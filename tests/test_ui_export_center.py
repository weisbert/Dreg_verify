# -*- coding: utf-8 -*-
"""test_ui_export_center.py —— ⑪ 导出中心 + ⑫ 导出完成（`dreg_verify/ui/export_center.py`，C4-a）。

契约 ID（附录 A · EXPORT 49 条）：C-157…C-199 + C-231 C-232 C-247 C-250 C-270 C-285
（C-172「预览完状态栏报数」归 `.sv` 预览那一路，已在 `tests/test_ui_sv_preview.py` 覆盖）。
不变量：I-03（导出选项单一来源）、I-09（配置字段集冻结 + 摘要只渲染一次）、I-20（名字在计数前）。

这一整个文件守的是同一件事：**产物和产物的账目必须对得上**。红区那位工程师拿到的是一个 .sv
和一句「7 信号 · 168 用例」——如果有两个信号被静默跳过、或者选项没落到 `exports` 上，
他要等到仿真跑完才发现，而那时表已经交出去了。所以每条都用真 mirror + 真 `WorkbenchState`
走真流程，并与 v1（`gui.MainWindow` / `SignalView`）逐字节对照。

夹具一律用 mirror 生成脚本（公开仓，**无真实信号名**）。
"""

import io
import json
import os
import zipfile

import pytest

import ui_harness as H                                    # noqa: E402

pytest.importorskip("PySide6")

from PySide6 import QtCore, QtWidgets                     # noqa: E402

from dreg_verify import edits as ED                       # noqa: E402
from dreg_verify import exports as X                      # noqa: E402
from dreg_verify import session                           # noqa: E402
from dreg_verify.ui import contracts                      # noqa: E402
from dreg_verify.ui import export_center as EC            # noqa: E402
from dreg_verify.ui import names                          # noqa: E402
from dreg_verify.ui import persist as P                   # noqa: E402
from dreg_verify.ui import state as ST                    # noqa: E402
from dreg_verify.ui import terms                          # noqa: E402

#: btlp 镜像里的一个 RO 回读根 —— 它**按设计**不产断言（`accounted`，C-166 点的就是它）
RO_SIG = "pll_lock_indicator"


# ═══════════════════════════ 夹具 ═══════════════════════════
@pytest.fixture(scope="module")
def qapp():
    return H.app()


@pytest.fixture(scope="module")
def btlp():
    return H.mirror_path("btlp")


@pytest.fixture
def iso(monkeypatch, tmp_path):
    """两份持久化文件指到 tmp（与 `test_ui_state.py` 的 `iso` 同一件事）。"""
    monkeypatch.setattr(P, "SETTINGS_PATH", str(tmp_path / "gui_settings.json"))
    monkeypatch.setattr(P, "EDITS_PATH", str(tmp_path / "edits.json"))
    return tmp_path


@pytest.fixture
def st(qapp, iso, btlp):
    """真 `WorkbenchState` + 骨架清单（= app 在 worker 起跑前的那两步）。"""
    s = ST.WorkbenchState()
    assert s.load(btlp)
    s.set_models("topout", s.provider("topout").skeleton_models(), partial=True)
    return s


@pytest.fixture
def rec(monkeypatch, tmp_path):
    """所有模态入口拦截 + 落盘目录（导出一路跑到底并真写文件）。"""
    d = tmp_path / "save"
    d.mkdir()
    return H.auto_dialogs(monkeypatch, save_dir=str(d))


def _enable(rows, *kinds):
    for r in rows:
        r.enabled = r.kind in kinds
    return rows


def _run(st, rows, save_dir, path_of=None):
    """按 rows 跑一趟（`ask_path` 默认落到 save_dir/默认文件名）。→ ExportRunResult。"""
    plan = EC.build_plan(st, rows)
    ask = path_of or (lambda row: os.path.join(str(save_dir), EC.default_filename(row, st)))
    return EC.run_plan(st, plan, ask)


def _dlg(st, **kw):
    d = EC.ExportCenterDialog(st, **kw)
    d.resize(1080, 520)
    return d


# ═══════════════ C-157…C-161 · C-163：.sv 的四个选项真的落到 exports ═══════════════
@pytest.mark.contract("C-157", "C-158", "C-159", "C-160", "C-161", "C-163")
def test_c157_c158_c159_c160_c161_sv_row_options_to_exports(st, rec, tmp_path, monkeypatch):
    """勾选框上的三个选项 + 范围下拉，一字不差地变成 `exports.export_sv(options=...)`。

    这四样**改产物**（注释 / 末尾汇总块 / 断言消息尾巴 / 正负向范围）。界面上勾了、
    产物里没有，用户是看不出来的 —— 所以这里盯的是「选项 → 后端参数」这一跳本身。
    """
    seen = {}
    real = X.export_sv

    def spy(provider, path, **kw):
        seen.update(kw)
        return real(provider, path, **kw)

    monkeypatch.setattr(X, "export_sv", spy)

    d = _dlg(st)
    d.show()
    # 三个选项全勾上（C-159 / C-160 / C-161）
    d.open_options("sv")
    for key in ("comments", "sv_summary", "owner_in_msg"):
        cb = H.find(d.options_pop, EC.fmt_export_opt("sv", key))
        cb.setChecked(True)
    d.options_pop.close_popover()
    # 范围切「仅正向」（C-158）
    sv_combo = H.find(d, names.fmt_export_scope_sv("sv"))
    sv_combo.setCurrentIndex(sv_combo.findData("pos"))
    assert d.row_of("sv").sv_scope == "pos"

    _enable(d.rows(), "sv")
    for k, cb in d.checks.items():
        cb.setChecked(k == "sv")
    out = d.run()

    assert seen.get("options") == {"comments": True, "sv_summary": True,
                                   "owner_in_msg": True, "scope": "pos"}
    assert len(out.outcomes) == 1 and out.outcomes[0].kind == "sv"
    text = io.open(out.outcomes[0].path, encoding="utf-8").read()
    assert text.startswith("// auto-generated")            # C-159 注释真的进了产物
    assert "dreg_n_real_fail" in text                      # C-160 末尾汇总计数器块
    # C-158：仅正向 → 产物里没有 _NEG 用例
    assert "_NEG" not in text
    d.close()


# ═══════════════ C-169 / C-170：分文件导出写两份、默认名带 _pos/_neg ═══════════════
@pytest.mark.contract("C-169", "C-170")
def test_c169_c170_split_writes_two_files(st, tmp_path):
    """「正向+反例分文件」= 两次 `export_sv`，两份文件、汇总命名块各带 `_pos`/`_neg` 后缀。

    后缀不是装饰：两份贴进同一个 testcase 作用域时，同名的命名 begin/end 块是非法 SV。
    """
    rows = _enable(EC.default_rows(st), "sv")
    row = next(r for r in rows if r.kind == "sv")
    row.options = {"comments": False, "sv_summary": True, "owner_in_msg": False}

    # C-169：三种 scope 的默认文件名各不相同
    row.sv_scope = "all"
    assert EC.default_filename(row, st) == terms.EXPORT_DEFAULT_SV
    row.sv_scope = "pos"
    assert EC.default_filename(row, st) == terms.EXPORT_DEFAULT_SV_POS
    row.sv_scope = "neg"
    assert EC.default_filename(row, st) == terms.EXPORT_DEFAULT_SV_NEG

    row.sv_scope = "split"
    base = str(tmp_path / "wr_rf_tc.sv")
    res = _run(st, rows, tmp_path, path_of=lambda _r: base)
    pos_path, neg_path = X.split_sv_paths(base)
    assert [o.path for o in res.outcomes] == [pos_path, neg_path]
    assert os.path.isfile(pos_path) and os.path.isfile(neg_path)
    pos = io.open(pos_path, encoding="utf-8").read()
    neg = io.open(neg_path, encoding="utf-8").read()
    # C-170：汇总命名块（`begin : xxx`）各带 _pos / _neg 后缀 —— 两份贴进同一作用域不重名
    blocks_pos = {ln.split(":", 1)[1].strip() for ln in pos.splitlines() if "begin :" in ln}
    blocks_neg = {ln.split(":", 1)[1].strip() for ln in neg.splitlines() if "begin :" in ln}
    assert blocks_pos and blocks_neg
    assert all(b.endswith("_pos") for b in blocks_pos), blocks_pos
    assert all(b.endswith("_neg") for b in blocks_neg), blocks_neg
    assert not (blocks_pos & blocks_neg)


# ═══════════════ C-164：重复 assert 标号 —— 不确认就一个字节都不写 ═══════════════
@pytest.mark.contract("C-164")
def test_c164_dup_labels_confirm_before_write(st, rec, tmp_path, monkeypatch):
    """重复 assert 标号 = elaboration 必失败的非法 SV。写文件**前**弹确认，按取消则不写。"""
    real_plan = EC.build_plan

    def with_dups(state, rows):
        plan = real_plan(state, rows)
        plan.dup_labels = [("R_DUP", "sig_a", "sig_b")]
        return plan

    monkeypatch.setattr(EC, "build_plan", with_dups)

    written = []
    monkeypatch.setattr(X, "export_sv",
                        lambda *a, **k: written.append(a[1]) or (_ for _ in ()).throw(
                            AssertionError("取消了还写文件")))

    # C4-int：确认框的答案走 harness（`DupLabelsDialog.ask` 已在 `CUSTOM_MODALS` 里），
    # 本文件不再自己 monkeypatch 对话框类 —— 同一个框两处各拦一套，改了一处就会各说各话。
    rec2 = H.auto_dialogs(monkeypatch, save_dir=rec.save_dir,
                          answers={"DupLabelsDialog.ask": False})

    d = _dlg(st)
    for k, cb in d.checks.items():
        cb.setChecked(k == "sv")
    assert d.run() is None, "按了取消还是导出了"
    asked = [list(c.args[1]) for c in rec2.of("DupLabelsDialog.ask")]   # args[0] = cls
    assert asked == [[("R_DUP", "sig_a", "sig_b")]], "没弹确认或没把冲突标号列出来"
    assert not written
    assert not [f for f in os.listdir(rec.save_dir) if f.endswith(".sv")]
    d.close()


# ═══════════════ C-165 / C-166 / C-167 / C-199：完成弹层的版式就是规范 ═══════════════
@pytest.mark.contract("C-165", "C-166", "C-167", "C-199", "C-270")
def test_c165_c166_c167_c199_done_dialog_order(st, qapp, tmp_path):
    """琥珀的「跳过点名」块在**上**、「已写出」计数块在**下**；名字在计数前。

    这条顺序是 V2Spec §6 那条文案红线在界面上的样子。反过来排（先给一个大数字再让人翻名字）
    是本项目反复纠正过的老毛病 —— 所以它有一条按**坐标**断言的测试，不只是断言文本存在。
    """
    rows = _enable(EC.default_rows(st), "sv")
    res = _run(st, rows, tmp_path)
    assert res.skipped, "这张镜像表本该有跳过项，测试样本选错了"

    d = EC.ExportDoneDialog(res)
    d.resize(820, 460)
    d.show()
    qapp.processEvents()

    sk = H.find(d, names.DONE_SKIPPED_BLOCK)
    wr = H.find(d, names.DONE_WRITTEN_BLOCK)
    y_sk = sk.mapTo(d, QtCore.QPoint(0, 0)).y()
    y_wr = wr.mapTo(d, QtCore.QPoint(0, 0)).y()
    assert y_sk < y_wr, "跳过点名块被排到了「已写出」下面（y %d ≥ %d）" % (y_sk, y_wr)

    # C-167：点名 + 原因，可展开明细
    body = d.skipped_text(full=True)
    for name, _why in res.skipped:
        assert name in body
    assert terms.DONE_SKIPPED_ROW_FMT.format(reason="").strip()[:1] in body or "└" in body
    short = d.skipped_text(full=False)
    d._toggle_detail()
    assert len(d.skipped_text()) >= len(short)

    # C-165：已写出那一行 = 路径 + 计数（信号 / 断言块 / 用例 / 反例 / 手填期望 / 兜底）
    line = d.written_lines()[0]
    assert res.outcomes[0].path in line
    for label in ("信号", "断言块", "测试用例"):
        assert label in line

    # C-166：点名「只记录、不产生断言」的信号（不是只给个数）
    assert RO_SIG in res.accounted
    assert RO_SIG in d.accounted_label.text()
    assert d.accounted_label.text().startswith(terms.DONE_ACCOUNTED_FMT.format(names="")[:6])

    # C-199 / I-20：整段文本里，第一个被点名的信号出现在「已写出」之前
    txt = d.text()
    assert txt.index(res.skipped[0][0]) < txt.index(terms.DONE_WRITTEN)
    d.close()


# ═══════════════ C-168：完成弹层关掉后切 .sv 预览 ═══════════════
@pytest.mark.contract("C-168")
def test_c168_sv_preview_requested_after_done(st, rec, qapp):
    """导出了 .sv → 关掉完成弹层后请求切到 .sv 预览；没导 .sv 就不该请求。"""
    d = _dlg(st)
    got = []
    d.svPreviewRequested.connect(lambda: got.append(1))
    for k, cb in d.checks.items():
        cb.setChecked(k == "sv")
    d.run()
    assert got == [1]
    d.close()

    d2 = _dlg(st)
    got2 = []
    d2.svPreviewRequested.connect(lambda: got2.append(1))
    for k, cb in d2.checks.items():
        cb.setChecked(k == "claims")
    d2.run()
    assert got2 == [], "没导 .sv 也去切 .sv 预览标签"
    d2.close()


# ═══════════════ C-171：写文件失败要说清是不是被占用 ═══════════════
@pytest.mark.contract("C-171")
def test_c171_export_error_message_names_lock(st, qapp, tmp_path):
    """.sv 正被仿真器/编辑器占用时写不进去 —— 弹层必须把这个可能性说出来，不是「导出失败」四个字。"""
    rows = _enable(EC.default_rows(st), "sv")
    bad = str(tmp_path / "no_such_dir" / "wr_rf_tc.sv")       # 目录不存在 → 真 OSError
    res = _run(st, rows, tmp_path, path_of=lambda _r: bad)
    assert not res.outcomes
    assert len(res.errors) == 1 and res.errors[0][0] == "sv"
    msg = res.errors[0][1]
    assert "占用" in msg and bad in msg

    d = EC.ExportDoneDialog(res)
    d.show()
    qapp.processEvents()
    err = H.find(d, names.DONE_ERRORS)
    assert err.isVisibleTo(d) and "占用" in err.text()
    assert terms.EXPORT_ROWS["sv"][0] in err.text()            # 哪一种交付物失败了
    d.close()


# ═══════════════ C-173 / C-174 / C-175 / C-176：报告三格式 + 扩展名纠正 ═══════════════
@pytest.mark.contract("C-173", "C-174", "C-175", "C-176")
def test_c173_c174_c175_c176_report_formats_and_ext(st, qapp, tmp_path):
    """三种格式各出各的文件；**选了 Excel 就存成 .xlsx**，不出「用例表.html.xlsx」这种双扩展名。"""
    d = _dlg(st)
    d.show()
    d.open_options("report")
    for key in contracts.EXPORT_REPORT_FORMATS:               # 三个单选都在（C-173/174/175）
        assert H.find(d.options_pop, EC.fmt_export_opt("report", key)) is not None
    d.options_pop.close_popover()

    row = d.row_of("report")
    for fmt, ext in (("html", ".html"), ("csv", ".csv"), ("xlsx", ".xlsx")):
        d._set_report_format(fmt)
        assert EC.default_filename(row, st).endswith(ext)
        assert EC.report_filter(fmt) in X.REPORT_FILTER_STR

    # C-176：默认名还是 .html、格式选了 Excel → 落盘必须是 .xlsx（`correct_report_ext` 权威）
    d._set_report_format("xlsx")
    rows = _enable(d.rows(), "report")
    wrong = str(tmp_path / "用例表.html")
    res = _run(st, rows, tmp_path, path_of=lambda _r: wrong)
    assert len(res.outcomes) == 1
    got = res.outcomes[0].path
    assert got.endswith(".xlsx") and not got.endswith(".html.xlsx")
    assert os.path.isfile(got)
    assert res.outcomes[0].kind == "report_xlsx"
    d.close()


# ═══════════════ C-177 / C-178：报告只导勾选的信号 + 完成报范围/用例/反例 ═══════════════
@pytest.mark.contract("C-177", "C-178")
def test_c177_c178_report_scope_and_summary(st, qapp, tmp_path, monkeypatch):
    """「范围」列不是摆设：切成勾选 → `render_report(only=...)` 真的只拿那几个。"""
    seen = {}
    real = st.provider("topout").render_report
    monkeypatch.setattr(st.provider("topout"), "render_report",
                        lambda *a, **k: seen.update(k) or real(*a, **k))

    keep = [m["name"] for m in st.models()][:3]
    st.set_checked([m["name"] for m in st.models()], False)
    st.set_checked(keep, True)

    rows = _enable(EC.default_rows(st), "report")
    row = next(r for r in rows if r.kind == "report")
    row.options = {"format": "csv"}
    row.object_scope = "checked"
    res = _run(st, rows, tmp_path)
    assert sorted(seen["only"]) == sorted(keep)                # C-178

    out = res.outcomes[0]
    detail = getattr(out, "ui_detail", "")
    assert detail == terms.EXPORT_REPORT_DONE_FMT.format(      # C-177
        scope=terms.EXPORT_SCOPE_CHECKED_FMT.format(n=len(keep)),
        n=out.counts["n_tests"], neg=out.counts["n_neg"])
    assert len(out.paths) == 3                                 # CSV 一次写三份扁平表，逐份列出
    d = EC.ExportDoneDialog(res)
    lines = "\n".join(d.written_lines())
    for p in out.paths:
        assert p in lines

    # 切回「全部」→ only=None（引擎的「不过滤」）
    row.object_scope = "all"
    _run(st, rows, tmp_path)
    assert seen["only"] is None


# ═══════════════ C-179…C-184：回填 for_test ═══════════════
@pytest.mark.contract("C-179", "C-180", "C-181", "C-182", "C-183", "C-184")
def test_c179_c180_c181_c182_c183_c184_fortest(st, qapp, tmp_path, monkeypatch, btlp):
    """回填是**写副本**：源表一个字节不动，扩展名自动补，回填了几组要报出来。"""
    rows = _enable(EC.default_rows(st), "fortest")
    row = next(r for r in rows if r.kind == "fortest")

    seen = {}
    real = X.export_fortest
    monkeypatch.setattr(X, "export_fortest",
                        lambda src, path, **k: seen.update(dict(k, src=src)) or real(src, path, **k))

    # C-183：只导勾选的
    keep = [m["name"] for m in st.models()][:2]
    st.set_checked([m["name"] for m in st.models()], False)
    st.set_checked(keep, True)
    row.object_scope = "checked"

    # C-181：给一个没有扩展名的路径 → 自动补 .xlsx
    noext = str(tmp_path / "backfill")
    before = os.path.getmtime(btlp), os.path.getsize(btlp)
    res = _run(st, rows, tmp_path, path_of=lambda _r: noext)
    assert not res.errors, res.errors
    out = res.outcomes[0]
    assert out.path == noext + ".xlsx" and os.path.isfile(out.path)     # C-179 / C-181
    assert sorted(seen["only"]) == sorted(keep)                         # C-183
    assert seen["provider"] is not None                                 # C-184：provider 路径含 mux 表
    assert out.note == "含 mux 表"
    assert (os.path.getmtime(btlp), os.path.getsize(btlp)) == before, "源表被改了"

    # C-182：完成报组数
    assert getattr(out, "ui_detail", "") == terms.EXPORT_FORTEST_DONE_FMT.format(
        n=out.counts["n_groups"])

    # C-180：拒绝把源 Excel 当输出名
    res2 = _run(st, rows, tmp_path, path_of=lambda _r: btlp)
    assert not res2.outcomes and res2.errors
    assert "源" in res2.errors[0][1]


# ═══════════════ C-185…C-189：nets.txt 按用途 + 按页 ═══════════════
@pytest.mark.contract("C-185", "C-186", "C-187", "C-188", "C-189")
def test_c185_c186_c187_c188_c189_nets_purpose_and_pages(st, rec, qapp, tmp_path, iso):
    """三用途默认全勾（拍板 #9）；「更多」按页只列**本表真有内容**的页；勾选记住上次。"""
    d = _dlg(st)
    d.show()
    row = d.row_of("nets")
    assert list(row.options["purposes"]) == list(EC.NETS_PURPOSES)      # C-186 默认全勾

    d.open_options("nets")
    for key in EC.NETS_PURPOSES:
        assert H.find(d.options_pop, EC.fmt_export_opt("nets", key)).isChecked()
    cats = EC.nets_categories(st)
    assert cats == list(X.nets_categories(st.wb))
    assert "iddq" not in cats, "这张镜像表没有 iddq 页，类别里不该出现"   # C-187
    for page in cats:
        assert H.find(d.options_pop, EC.fmt_export_nets_page(page)) is not None
    # C-188 / C-232：勾一页 → 记进 settings 的 nets_pages
    H.find(d.options_pop, EC.fmt_export_nets_page("logic")).setChecked(True)
    assert st.settings()["nets_pages"] == ["logic"]
    d.options_pop.close_popover()

    # 新开一个导出中心 → 预选上次那一页
    d2 = _dlg(st)
    assert d2.row_of("nets").options["pages"] == ["logic"]
    d2.close()

    # C-185 / C-189：真导一次 —— 总网数 + 各类别分项计数
    for k, cb in d.checks.items():
        cb.setChecked(k == "nets")
    res = d.run()
    out = res.outcomes[0]
    assert out.path.endswith(".txt") and os.path.isfile(out.path)
    assert out.counts["n_nets"] > 0
    for purpose in EC.NETS_PURPOSES:
        assert purpose in out.counts                                    # C-189 分项计数
    assert out.counts["logic"] >= 0                                     # 「更多」那一页也计了数
    assert list(out.counts)[0] == "n_nets"
    d.close()


# ═══════════════ C-196：claims.json ═══════════════
@pytest.mark.contract("C-196")
def test_c196_claims_row(st, qapp, tmp_path):
    """第 5 行 claims.json —— v1 GUI 根本没有这个出口，v2 必须有，且与 CLI 的写出器同构。"""
    rows = _enable(EC.default_rows(st), "claims")
    row = next(r for r in rows if r.kind == "claims")
    assert EC.default_filename(row, st) == terms.EXPORT_DEFAULT_CLAIMS
    res = _run(st, rows, tmp_path)
    out = res.outcomes[0]
    payload = json.load(io.open(out.path, encoding="utf-8"))
    assert payload["claims"] and payload.get("naming_model") == "topout"
    kinds = {c["kind"] for c in payload["claims"]}
    assert {"probe", "force"} <= kinds
    for key in ("n_claims", "n_probe", "n_force", "n_rfwrite"):
        assert key in out.counts
    assert st.last_export("claims").path == out.path


# ═══════════════ C-197 / C-198：一张表 6 行 + 「上次导出到哪」列 ═══════════════
@pytest.mark.contract("C-197", "C-198")
def test_c197_c198_six_rows_last_export_column(st, rec, qapp):
    """6 种交付物 × 5 列一张表；每种记住上次导出到哪（「我到底导没导过 nets.txt」）。"""
    d = _dlg(st)
    d.show()
    qapp.processEvents()
    t = H.find(d, names.EXPORT_TABLE)
    assert t.rowCount() == 6 and t.columnCount() == 5
    assert H.header_texts(t) == list(terms.EXPORT_HEADERS)
    assert [r.kind for r in d.rows()] == list(contracts.EXPORT_KINDS)
    # 默认勾 .sv / 报告 / claims.json（Design ER 的 [0,1,4]）
    assert [r.kind for r in d.rows() if r.enabled] == ["sv", "report", "claims"]
    for kind in contracts.EXPORT_KINDS:                      # 五列每行都在
        for fmt in (names.fmt_export_check, names.fmt_export_scope,
                    names.fmt_export_scope_objects, names.fmt_export_options,
                    names.fmt_export_last):
            assert H.find(d, fmt(kind)) is not None
        assert H.find(d, names.fmt_export_last(kind)).text() == terms.EXPORT_LAST_NEVER
    assert H.find(d, names.fmt_export_scope_objects("config")).text() == terms.EXPORT_SCOPE_NA
    assert H.find(d, names.fmt_export_scope_sv("sv")) is not None

    for k, cb in d.checks.items():
        cb.setChecked(k == "nets")
    res = d.run()
    d.close()

    le = st.last_export("nets")                              # C-198：落进 settings
    assert le is not None and le.path == res.outcomes[0].path
    d2 = _dlg(st)
    assert H.find(d2, names.fmt_export_last("nets")).text() == EC.last_text(le)
    assert le.path in H.find(d2, names.fmt_export_last("nets")).text()
    assert H.find(d2, names.fmt_export_last("sv")).text() == terms.EXPORT_LAST_NEVER
    d2.close()


# ═══════════════ C-190…C-195 / C-247 / C-250：配置往返 ═══════════════
@pytest.mark.contract("C-190", "C-191", "C-192", "C-193", "C-194", "C-195", "C-247", "C-250")
def test_c190_c191_c192_c193_c194_c195_c247_c250_config_roundtrip(st, qapp, tmp_path, btlp):
    """导出整份配置 → 改表名导入 → 三段提示；字段集与 v1 逐键相同（老同事的文件仍能导入）。"""
    st.set_probe_prefixes({"d_probe_demo": "U_TOP.U_SUB"})
    st.set_force_signals({"d_force_demo"})
    keep = [m["name"] for m in st.models()][:4]
    st.set_checked([m["name"] for m in st.models()], False)
    st.set_checked(keep, True)

    rows = _enable(EC.default_rows(st), "config")
    row = next(r for r in rows if r.kind == "config")
    assert EC.default_filename(row, st) == session.default_config_filename(btlp)   # C-190
    res = _run(st, rows, tmp_path)
    cfg_path = res.outcomes[0].path
    payload = json.load(io.open(cfg_path, encoding="utf-8"))

    # C-247：版本号；C-250：字段集与 v1 `_collect_config()` 逐键相同
    assert payload["dreg_verify_config"] == session.CONFIG_VERSION == 2
    v1 = session.collect_config("x.xlsx", {"coverage_logic": "全面", "coverage_mux": "全面",
                                           "max_tests": 256, "cascade_logic": "cone",
                                           "cascade_mux": "cone", "append_to_logic": True,
                                           "append_to_mux": False, "include_risky": True})
    assert sorted(payload) == sorted(v1)
    assert sorted(payload["global"]) == sorted(v1["global"])
    assert sorted(payload["signals_checked"]) == sorted(keep)
    assert payload["probe_prefixes"] == {"d_probe_demo": "U_TOP.U_SUB"}
    assert payload["force_signals"] == ["d_force_demo"]

    # C-191：导出后逐段报数
    detail = getattr(res.outcomes[0], "ui_detail", "")
    assert detail.startswith("配置已导出：") and "勾选 4 个" in detail and "前缀 1 条" in detail

    # C-195：不是本工具的配置 → 说清缺哪个段
    junk = tmp_path / "junk.json"
    junk.write_text(json.dumps({"hello": 1}), encoding="utf-8")
    bad = EC.import_config(st, str(junk))
    assert not bad.ok and bad.error == terms.EXPORT_IMPORT_BAD_FILE

    # C-193（表名核对）+ C-194（点名找不到的信号，名字在计数前）
    payload["excel"] = "另一张表.xlsx"
    payload["view_edits"] = {"topout": {"no_such_signal_x": {"cols": []},
                                        "no_such_signal_y": {"cols": []}}}
    io.open(cfg_path, "w", encoding="utf-8").write(json.dumps(payload, ensure_ascii=False))
    rep = EC.import_config(st, cfg_path)
    assert rep.ok and rep.is_full                                        # C-192
    assert any(terms.EXPORT_IMPORT_MISMATCH_FMT.format(cfg="另一张表.xlsx",
                                                       cur=os.path.basename(btlp)) == n
               for n in rep.notes)
    assert sorted(n for n, _r in rep.missing) == ["no_such_signal_x", "no_such_signal_y"]
    txt = rep.text()
    assert txt.index("no_such_signal_x") < txt.index("共 2 个")           # C-194 / I-20
    # 导入后勾选照单恢复（不是全勾）
    assert sorted(st.checked_names()) == sorted(keep)
    assert st.probe_prefixes == {"d_probe_demo": "U_TOP.U_SUB"}

    # C-192 的另一半：旧版【测试项编辑】文件 = 只并入编辑，别的一律不动
    legacy = tmp_path / "legacy.json"
    legacy.write_text(json.dumps({"edits": {}, "mux_expected": {}}), encoding="utf-8")
    lrep = EC.import_config(st, str(legacy))
    assert lrep.ok and lrep.is_legacy and not lrep.is_full


# ═══════════════ C-256：Ctrl+R 打开导出中心并只勾报告行 ═══════════════
@pytest.mark.contract("C-256")
def test_c256_ctrl_r_preselects_report(st, qapp):
    """Ctrl+R 不是「直接导报告」，是「打开导出中心并预选报告行」——路径、格式仍由人定。"""
    assert contracts.SHORTCUTS["export_report"] == "Ctrl+R"
    d = _dlg(st, preselect="report")
    assert [r.kind for r in d.rows() if r.enabled] == ["report"]
    assert H.find(d, names.fmt_export_check("report")).isChecked()
    assert not H.find(d, names.fmt_export_check("sv")).isChecked()
    assert d.run_btn.text() == terms.EXPORT_BTN_RUN_FMT.format(k=1)
    d.close()

    d2 = _dlg(st, preselect="nets")                          # 诊断抽屉第 1 步
    assert [r.kind for r in d2.rows() if r.enabled] == ["nets"]
    d2.close()

    d3 = _dlg(st)                                            # 不预选 → Design 的 [0,1,4]
    assert [r.kind for r in d3.rows() if r.enabled] == ["sv", "report", "claims"]
    d3.close()


# ═══════════════ C-162 / C-231 / C-232：选项记住上次 ═══════════════
@pytest.mark.contract("C-162", "C-231", "C-232")
def test_c162_c231_c232_options_remembered(st, qapp, iso):
    """四个 .sv 选项 + nets 按页类别记住上次选择，下次预选。"""
    d = _dlg(st)
    d.show()
    d.open_options("sv")
    H.find(d.options_pop, EC.fmt_export_opt("sv", "comments")).setChecked(True)
    H.find(d.options_pop, EC.fmt_export_opt("sv", "owner_in_msg")).setChecked(True)
    d.options_pop.close_popover()
    sv_combo = H.find(d, names.fmt_export_scope_sv("sv"))
    sv_combo.setCurrentIndex(sv_combo.findData("neg"))
    d.close()

    # 四个 export_* 键都落了盘（C-231），且读回来还是同一套（C-162）
    opt = P.load_export_options()
    assert opt == {"scope": "neg", "comments": True, "sv_summary": False, "owner_in_msg": True}
    st_file = json.load(io.open(P.SETTINGS_PATH, encoding="utf-8"))
    assert set(X.EXPORT_OPTION_KEYS.values()) <= set(st_file)

    d2 = _dlg(st)
    row = d2.row_of("sv")
    assert row.sv_scope == "neg"
    assert row.options == {"comments": True, "sv_summary": False, "owner_in_msg": True}
    d2.show()
    d2.open_options("sv")
    assert H.find(d2.options_pop, EC.fmt_export_opt("sv", "comments")).isChecked()
    assert not H.find(d2.options_pop, EC.fmt_export_opt("sv", "sv_summary")).isChecked()
    assert H.find(d2, names.fmt_export_scope_sv("sv")).currentData() == "neg"
    d2.close()


# ═══════════════ C-285：HTML 报告内联电路图 ═══════════════
@pytest.mark.contract("C-285")
def test_c285_report_html_inlines_flow_svg(st, qapp, tmp_path):
    """HTML 报告里每个信号带一张内联电路图 —— 「这张表在验的到底是个什么电路」当场看得见。

    走 `exports.attach_report_graphs`（= CLI 默认那条路），视图层不 import sigflow（I-19）。
    CSV / Excel 两种格式不该为此多跑一趟建图。
    """
    rows = _enable(EC.default_rows(st), "report")
    row = next(r for r in rows if r.kind == "report")
    row.options = {"format": "html"}
    res = _run(st, rows, tmp_path)
    html = io.open(res.outcomes[0].path, encoding="utf-8").read()
    assert html.count("<svg") >= 3, "HTML 报告里一张电路图都没有"
    assert "sigflowbox" in html

    calls = []
    real = X.attach_report_graphs
    try:
        X.attach_report_graphs = lambda *a, **k: calls.append(1) or real(*a, **k)
        row.options = {"format": "csv"}
        _run(st, rows, tmp_path)
        assert calls == [], "CSV 报告也去建了图（白跑一趟引擎）"
    finally:
        X.attach_report_graphs = real


# ═══════════════ I-03：导出选项只有一个来源 ═══════════════
@pytest.mark.contract("C-163", "C-231")
def test_i03_export_options_single_source(st, qapp, monkeypatch):
    """I-03：四个 export_* 键**只**经 `exports.load/store_export_options`。

    两个旧对话框各写一套默认值、同一份磁盘配置被两套口径解释，正是 C-163 要合掉的那个坑；
    本模块里再出现一次 `"export_scope"` 这样的字面量，就是坑回来了。
    """
    src = io.open(EC.__file__, encoding="utf-8").read()
    for key in X.EXPORT_OPTION_KEYS.values():
        assert '"%s"' % key not in src and "'%s'" % key not in src, \
            "export_center.py 里直接写了 settings 键 %r —— 要走 exports 的 load/store" % key

    loads, stores = [], []
    monkeypatch.setattr(P, "load_export_options",
                        lambda: loads.append(1) or dict(X.EXPORT_OPTION_DEFAULTS))
    monkeypatch.setattr(P, "store_export_options", lambda opt: stores.append(dict(opt)))
    d = _dlg(st)
    assert loads, "六行默认值不是从 exports 的 load_export_options 来的"
    assert d.row_of("sv").options == {k: X.EXPORT_OPTION_DEFAULTS[k]
                                      for k in ("comments", "sv_summary", "owner_in_msg")}
    d._set_sv_option("sv_summary", True)
    assert stores and stores[-1]["sv_summary"] is True and "scope" in stores[-1]
    d.close()


# ═══════════════ I-09：导出前摘要只渲染一次 ═══════════════
def test_i09_build_plan_renders_once(st, qapp, monkeypatch):
    """I-09：`build_plan` 整个只渲染**一次** .sv。

    六行里 .sv / claims / 报告都要同一趟展开；每行各渲一次 = 200 个信号的展开跑三遍。
    这条测试是那个「跑三遍」不会悄悄回来的唯一保证。
    """
    n_x, n_prov = [], []
    real_x = X.render_sv
    monkeypatch.setattr(X, "render_sv", lambda *a, **k: n_x.append(1) or real_x(*a, **k))
    prov = st.provider("topout")
    real_p = prov.render_sv
    monkeypatch.setattr(prov, "render_sv", lambda *a, **k: n_prov.append(1) or real_p(*a, **k))

    rows = EC.default_rows(st)
    for r in rows:
        r.enabled = True                       # 六行全勾，最坏情况
    plan = EC.build_plan(st, rows)
    assert len(n_x) == 1, "build_plan 渲染了 %d 次 .sv" % len(n_x)
    assert len(n_prov) == 1
    assert plan.n_signals == len(st.models())

    # 一行都没勾 → 一次都不渲染（打开导出中心本身不该跑引擎）
    del n_x[:], n_prov[:]
    for r in rows:
        r.enabled = False
    EC.build_plan(st, rows)
    assert n_x == [] and n_prov == []


def test_toggling_a_row_does_not_rerender(st, qapp, monkeypatch):
    """勾/取消一行交付物不该重跑一遍展开 —— 真表 200 个信号时那是每点一下卡几秒。

    「哪些信号会被跳过」只由覆盖度指纹 + 对象范围决定，与勾了哪几种产物无关。
    """
    n = []
    real = X.render_sv
    monkeypatch.setattr(X, "render_sv", lambda *a, **k: n.append(1) or real(*a, **k))
    d = _dlg(st)
    assert len(n) == 1                                   # 开框时算一次
    for kind in ("fortest", "nets", "config", "report"):
        d.checks[kind].setChecked(not d.checks[kind].isChecked())
    assert len(n) == 1, "点了 4 下勾选框，引擎跑了 %d 趟" % len(n)
    # 对象范围变了（勾选的 N 个 → 全部）就该重算 —— 跳过项真的会不一样
    assert EC.plan_scope_row(d.rows()).kind == "sv"
    d.obj_combos["sv"].setCurrentIndex(list(contracts.EXPORT_OBJECT_SCOPES).index("all"))
    assert len(n) == 2, "对象范围换了却没重算摘要"
    d.close()


def test_summary_says_so_when_it_cannot_be_computed(st, qapp, monkeypatch):
    """摘要算不出来时说实话 —— **绝不**退化成「0 个信号会被跳过」那种让人放心的假数字。"""
    monkeypatch.setattr(X, "render_sv",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("引擎炸了")))
    d = _dlg(st)
    txt = d.summary.text()
    assert "引擎炸了" in txt
    assert terms.EXPORT_SUMMARY_NO_SKIP_FMT.format(k=3, n=len(st.models())) not in txt
    assert d.run_btn.isEnabled(), "摘要算不出来不该把导出按钮也锁死"
    d.close()


# ═══════════════ I-20：每一段完成文案都是「名字在前、计数在后」 ═══════════════
@pytest.mark.contract("C-199", "C-270")
def test_i20_every_done_text_names_before_counts(st, qapp, tmp_path):
    """I-20：导出前摘要 · 完成弹层 · 导入配置结果 —— 三处的点名都排在计数前面。"""
    rows = _enable(EC.default_rows(st), "sv")
    res = _run(st, rows, tmp_path)
    assert res.skipped

    # ① 完成弹层（C-167 / C-199）
    done = EC.ExportDoneDialog(res)
    txt = done.text()
    name = res.skipped[0][0]
    assert txt.index(name) < txt.index(terms.DONE_WRITTEN)
    # 琥珀块里：信号名那一行在它的原因行之前
    body = done.skipped_text(full=True).splitlines()
    assert body[0] == name and body[1].lstrip().startswith("└")

    # ② 导出前摘要（裁决⑯：点名可展开，不是只给个数）
    d = _dlg(st)
    d.show()
    assert d.skipped_btn.isVisible()
    assert name in d.skipped_text()
    assert d.skipped_text().index(name) == 0                 # 名字在原因前
    d.close()

    # ③ 导入配置结果（C-194）
    rep = EC.ImportReport()
    rep.missing = [("sig_gone_a", "当前表里没有这个信号"), ("sig_gone_b", "当前表里没有这个信号")]
    rep.counts = ["已导入完整配置：恢复了 0 个信号的手填编辑"]
    t = rep.text()
    assert t.index("sig_gone_a") < t.index("共 2 个")


# ═══════════════ 「去处理这 N 个信号」 / 「打开输出目录」 ═══════════════
def test_done_dialog_go_fix_and_open_dir(st, qapp, tmp_path, monkeypatch):
    """N6：打开输出目录（纯 GUI 一行）；跳过项存在时才出现「去处理这 N 个信号」。"""
    rows = _enable(EC.default_rows(st), "sv")
    res = _run(st, rows, tmp_path)
    opened = []
    from PySide6 import QtGui
    monkeypatch.setattr(QtGui.QDesktopServices, "openUrl",
                        staticmethod(lambda url: opened.append(url.toLocalFile()) or True))
    d = EC.ExportDoneDialog(res)
    d.show()
    qapp.processEvents()
    assert d.open_out_dir()
    assert os.path.normcase(opened[-1].rstrip("/")) == os.path.normcase(str(tmp_path))

    got = []
    d.goFixRequested.connect(got.append)
    assert d.go_fix_btn.isVisible()
    assert d.go_fix_btn.text() == terms.DONE_BTN_GO_FIX_FMT.format(n=len(res.skipped))
    H.click(d.go_fix_btn)
    # 症状键必须是 app 路由认得的那一个（不是随手编的字符串）
    assert got == [EC.GO_FIX_SYMPTOM]
    assert EC.GO_FIX_SYMPTOM in terms.REASON_TARGETS
    d.close()

    # 一个都没跳过 → 那个按钮不该出现（没东西可处理还放个按钮＝噪音）
    empty = contracts.ExportRunResult(outcomes=list(res.outcomes), out_dir=str(tmp_path))
    d2 = EC.ExportDoneDialog(empty)
    d2.show()
    qapp.processEvents()
    assert not d2.go_fix_btn.isVisible()
    assert not H.find(d2, names.DONE_SKIPPED_BLOCK).isVisibleTo(d2)
    d2.close()


# ═══════════════ 导入配置入口（C-192 的按钮那一半）═══════════════
@pytest.mark.contract("C-192")
def test_import_config_button_and_preselect(st, qapp, tmp_path, monkeypatch):
    """「导入配置…」按钮在导出中心底部；`preselect="config"` 打开即进导入流程。"""
    d = _dlg(st)
    assert H.find(d, names.EXPORT_BTN_IMPORT_CONFIG).text() == terms.EXPORT_BTN_IMPORT_CONFIG
    d.close()

    # 文件框走 harness（不再本地 monkeypatch QFileDialog）：答「取消」= 导入流程起了但没选文件
    rec = H.auto_dialogs(monkeypatch, answers={"getOpenFileName": ("", "")})
    d2 = _dlg(st, preselect="config")
    d2.show()
    qapp.processEvents()
    assert rec.count("QFileDialog.getOpenFileName") == 1, "preselect='config' 没有直接进导入配置"
    d2.close()


# ═══════════════ 一个都没勾时：按钮置灰 + 文案变 ═══════════════
def test_run_button_reflects_selection(st, qapp):
    d = _dlg(st)
    for cb in d.checks.values():
        cb.setChecked(False)
    assert not d.run_btn.isEnabled()
    assert d.run_btn.text() == terms.EXPORT_BTN_RUN_NONE
    assert d.run() is None
    d.checks["sv"].setChecked(True)
    assert d.run_btn.isEnabled()
    assert d.run_btn.text() == terms.EXPORT_BTN_RUN_FMT.format(k=1)
    d.close()


# ═══════════════ 用户在文件框按取消 = 这一行不导，不是错误 ═══════════════
def test_cancel_one_row_skips_only_that_row(st, qapp, tmp_path):
    rows = _enable(EC.default_rows(st), "sv", "claims")

    def ask(row):
        return "" if row.kind == "sv" else os.path.join(str(tmp_path),
                                                        EC.default_filename(row, st))

    res = _run(st, rows, tmp_path, path_of=ask)
    assert [o.kind for o in res.outcomes] == ["claims"]
    assert not res.errors
    assert st.last_export("sv") is None


# ═══════════════ 截图（Design ⑪⑫ 两张场景）═══════════════
def test_shots_export_center_and_done(st, qapp, tmp_path):
    d = _dlg(st)
    d.show()
    qapp.processEvents()
    assert os.path.isfile(H.shot(d, "export_center"))
    rows = _enable(EC.default_rows(st), "sv", "claims", "config")
    res = _run(st, rows, tmp_path)
    done = EC.ExportDoneDialog(res)
    done.resize(820, 460)
    done.show()
    qapp.processEvents()
    assert os.path.isfile(H.shot(done, "export_done"))
    done.close()
    d.close()


# ═══════════════════════════ v1 对照：逐字节 ═══════════════════════════
def _v1_window(monkeypatch, tmp_path, btlp):
    """一台真 v1 `MainWindow`（`legacy_gui.py`），载同一张 mirror 表。"""
    H.isolate_settings(monkeypatch, tmp_path)
    w = H.make_window(btlp)
    assert w.wb is not None and getattr(w, "topout_view", None) is not None
    return w


def _strip_ts(text):
    """报告 HTML/CSV 里的时间戳行剔掉再比（两次导出的墙上时间必然不同）。"""
    return "\n".join(ln for ln in (text or "").splitlines()
                     if "生成时间" not in ln and "generated at" not in ln.lower())


def _xlsx_cells(path):
    """xlsx → {sheet: [[cell…]]}（zip 里有时间戳，逐字节比不稳，比内容）。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    return {ws.title: [[c for c in row] for row in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets}


@pytest.mark.contract("C-157", "C-173", "C-179", "C-185", "C-196", "C-190")
def test_v1_parity_six_deliverables(qapp, monkeypatch, tmp_path, btlp):
    """v1 对照：同一张表、同一套勾选与选项下，v2 `run_plan` 写出的产物与 v1 的对应 `on_export_*` 一致。

    .sv / nets.txt / claims.json **逐字节**相同；报告 CSV 剔掉时间戳行后逐字节相同；
    for_test 回填比单元格内容（xlsx 是 zip，成员时间戳每次不同）；配置 .json 比字段集
    （两个门面持有的状态本就不同，能冻结的是 schema —— C-247 / C-250 要的也正是这个）。
    """
    v1_dir = tmp_path / "v1"
    v2_dir = tmp_path / "v2"
    v1_dir.mkdir()
    v2_dir.mkdir()
    # 全部模态入口走 harness（C4-int：本文件不再自己 monkeypatch 任何对话框入口）。
    # `_confirm_dup_labels`（默认 True）/ `_ask_export_options`（默认 = 全默认选项）
    # 用 harness 的默认答案即可；只有两样要给：逐次「另存为」的落点、nets 的按页类别。
    v1_saves = []                       # 按调用顺序弹出的落点（一次「另存为」取一个）
    cats_box = []                       # v1 载表后才知道本表真有哪几页（w 在下一行才有）
    H.auto_dialogs(monkeypatch, answers={
        "getSaveFileName": lambda call, *a, **k: ((str(v1_saves.pop(0)), "") if v1_saves
                                                  else pytest.fail("v1 多弹了一次「另存为」")),
        "_ask_nets_pages": lambda call, *a, **k: list(cats_box[0]),
    })
    opts = dict(X.EXPORT_OPTION_DEFAULTS)      # = harness `_ask_export_options` 的默认答案
    w = _v1_window(monkeypatch, tmp_path, btlp)
    try:
        view = w.topout_view

        def save_to(p):
            v1_saves.append(p)

        # ── v1 的六份产物 ──
        save_to(v1_dir / "wr_rf_tc.sv")
        view.on_export_sv()
        save_to(v1_dir / "report.csv")
        view.on_export_report()
        save_to(v1_dir / "backfill.xlsx")
        view.on_fortest()
        cats = list(X.nets_categories(w.wb))
        cats_box.append(cats)
        save_to(v1_dir / "nets.txt")
        w.on_export_nets()
        save_to(v1_dir / "cfg.json")
        w.on_export_edits()
        v1_claims = str(v1_dir / "claims.json")            # v1 GUI 无出口 → 用它的 provider 走 exports
        X.export_claims(view.provider, v1_claims, w._loaded_excel_path or "",
                        only=view._checked_names() or None, mode=view._mode()[0],
                        max_tests=view._maxt(), exhaustive=view._mode()[1],
                        edited=view._compute_edited(),
                        sig_cov=view._sig_cov, form_cov=view._form_cov)
        v1_cfg = json.load(io.open(str(v1_dir / "cfg.json"), encoding="utf-8"))
    finally:
        w.close()
        qapp.processEvents()

    # ── v2 的六份产物（同一张表、默认全勾、默认覆盖度档）──
    s = ST.WorkbenchState()
    assert s.load(btlp)
    s.set_models("topout", s.provider("topout").skeleton_models(), partial=True)
    rows = EC.default_rows(s)
    for r in rows:
        r.enabled = True
        r.object_scope = "all"                              # v1 未勾任何行时 only=None
    next(r for r in rows if r.kind == "sv").options = dict(
        (k, opts[k]) for k in ("comments", "sv_summary", "owner_in_msg"))
    next(r for r in rows if r.kind == "sv").sv_scope = "all"
    next(r for r in rows if r.kind == "report").options = {"format": "csv"}
    next(r for r in rows if r.kind == "nets").options = {"purposes": [], "pages": cats}
    fname = {"sv": "wr_rf_tc.sv", "report": "report.csv", "fortest": "backfill.xlsx",
             "nets": "nets.txt", "claims": "claims.json", "config": "cfg.json"}
    res = EC.run_plan(s, EC.build_plan(s, rows),
                      lambda row: str(v2_dir / fname[row.kind]))
    assert not res.errors, res.errors
    got = {o.kind: o for o in res.outcomes}

    # ① .sv —— 逐字节。先证明比的不是两个空文件（「都没写出来」也会让 == 成立）
    v1_sv = io.open(str(v1_dir / "wr_rf_tc.sv"), "rb").read()
    assert len(v1_sv) > 2000 and b"assert (" in v1_sv, "v1 那份 .sv 是空的，这条对照等于没比"
    assert io.open(got["sv"].path, "rb").read() == v1_sv, ".sv 与 v1 不是逐字节相同"
    # ② nets.txt —— 逐字节（v2 的「按用途」层关掉，只走 v1 的按页口径）
    v1_nets = io.open(str(v1_dir / "nets.txt"), "rb").read()
    assert len(v1_nets.splitlines()) > 5
    assert io.open(got["nets"].path, "rb").read() == v1_nets
    # ③ claims.json —— 逐字节
    v1_cl = io.open(v1_claims, "rb").read()
    assert json.loads(v1_cl.decode("utf-8"))["claims"]
    assert io.open(got["claims"].path, "rb").read() == v1_cl
    # ④ 报告 CSV 三份 —— 剔时间戳行后逐字节
    v1_csv = sorted(f for f in os.listdir(str(v1_dir)) if f.endswith(".csv"))
    v2_csv = sorted(f for f in os.listdir(str(v2_dir)) if f.endswith(".csv"))
    assert v1_csv == v2_csv and len(v1_csv) == 3
    for fn in v1_csv:
        a = _strip_ts(io.open(str(v1_dir / fn), encoding="utf-8-sig").read())
        b = _strip_ts(io.open(str(v2_dir / fn), encoding="utf-8-sig").read())
        assert len(a.splitlines()) > 3, "v1 的 %s 是空表" % fn
        assert a == b, "报告 %s 与 v1 不一致" % fn
    # ⑤ for_test 回填 —— 逐单元格
    v1_cells = _xlsx_cells(str(v1_dir / "backfill.xlsx"))
    assert any(len(rows) > 3 for rows in v1_cells.values())
    assert _xlsx_cells(got["fortest"].path) == v1_cells
    assert zipfile.is_zipfile(got["fortest"].path)
    # ⑥ 配置 .json —— 字段集（C-247 / C-250）
    v2_cfg = json.load(io.open(got["config"].path, encoding="utf-8"))
    assert sorted(v2_cfg) == sorted(v1_cfg)
    assert sorted(v2_cfg["global"]) == sorted(v1_cfg["global"])
    assert v2_cfg["dreg_verify_config"] == v1_cfg["dreg_verify_config"]
    assert v2_cfg["excel"] == v1_cfg["excel"]


# ═══════════════ 附：`is_accounted` 的判据只有一份 ═══════════════
def test_accounted_classifier_uses_exports_labels():
    """C-166 的「只记录、不产生断言」判据来自 `exports.SKIP_STATUS_LABEL`，本模块不另抄一套说法。"""
    for status in EC.ACCOUNTED_STATUSES:
        assert EC.is_accounted(X.skip_reason(status, "随便什么原文"))
    for status in ("unresolved", "error", "needs-prefix", "spec-collision", "false-green"):
        assert not EC.is_accounted(X.skip_reason(status, ""))
    assert not EC.is_accounted(X.risky_reason([("A", "d_x", "探不到")]))


def test_serialize_view_edits_is_the_only_config_edit_shape():
    """配置里的 `view_edits` 段形状 = `edits.serialize_view_edits`（导出写它、导入按它还原）。"""
    assert callable(ED.serialize_view_edits) and callable(ED.restore_view_edits)
    assert EC.collect_config.__doc__
