# -*- coding: utf-8 -*-
"""truth/io.py —— 真值表的剪贴板 / 导入 / 导出（**Qt-free**）。

架构 §6.15，契约 TRUTH_IO 6 条（C-136/C-137/C-138/C-139/C-294/C-298）。

分工（护栏）：

* **本模块不落格、不碰撤销栈**。`paste_plan` 只算「哪一格该写什么」，谁去写、裹不裹成一步
  撤销是 `truth/model.py` 的事（C3-int 会把 `model.paste_tsv` 改成调这里的 `paste_plan`，
  两边规则一字不差由 `test_c294_paste_plan_equivalent_to_model_paste` 钉死）。
* **本模块不 import PySide6**。`indexes` 只当作「有 `.row()` / `.column()` 的对象」，
  model 只当作 `contracts.TruthModelProto` 的面来调（`row_kind` / `cell_state` / `col_state`
  / `cols` / `editable_kind` / `rowCount` / `columnCount` / `data`）。
* **读文件在这边、写格子在 model 那边**（C-298）：`read_expectations` 只把
  {列名: 值} 读出来，认不认得上这个列名是 `model.apply_expectations` 报。

导出 CSV 的两条口径（别混）：

1. **行**（第一列的信号/字段名）走 `csv_input_rows(an)` —— 与 v1 `SignalView.e_inputs`
   逐字节同（真名带位宽，不带真值表冻结列那个 `(角色 · 端口)` 后缀）。CSV 是给
   designer 拿去比对 / 回填的交换件，不是界面。
2. **列**（每列一条测试）：logic / register / dft 走 model 现在这张表（所见即所得）；
   **mux 走产物**（C-139）—— 列集 = `exports.render_sv(provider, only=[name], edited=…)`
   那次 build 里该信号实际渲染进 .sv 的向量，经 `exports.columns_from_mux_vectors` 折成列。
   mux 的列不能拿 model 的 `cols()`：整信号负向（`which=first` 追加一条）、负向去重、
   T 编号重排、iddq 自检拍都发生在 build 里，编辑器手上那份看不到，照它导出的 CSV 与
   .sv 对不上 —— 而 designer 拿 CSV 就是来核对 .sv 的。
"""

import csv
import os
from dataclasses import dataclass, field

from ... import exports as X
from ... import inputs_table as IT
from ... import truth_edit as TE
from .. import contracts
from .. import terms

__all__ = ["PastePlan", "PasteTally", "BAD_NAMES_MAX",
           "parse_tsv", "copy_tsv", "paste_plan", "paste_report_text",
           "plan_added_names", "read_expectations", "csv_input_rows", "signal_csv_columns",
           "export_signal_csv", "import_report_text", "COV_KEYS"]


#: ⚠ C3-int 起本模块**不再有 `PENDING_TERMS` 影子表**：文案只在 `ui/terms.py`。
#: 粘贴 / 导入期望的**结果说明整句**也收在本模块（`paste_report_text` /
#: `import_report_text`）—— `truth/model.py` 与 `truth/panel.py` 都调这一份，
#: 不再各拼各的（C3-int 去重；`test_c294_paste_plan_equivalent_to_model_paste`
#: 逐字比两边的报告）。

#: CSV / xlsx 里「期望」那一行的行名——按这个顺序找，第一个命中的算数。
#: 与 `exports.signal_csv_text` 写出来的行名对齐（导出的表原样改回来就能导入）。
EXP_ROW_PREFIXES = ("期望(进.sv)", "期望(bin)", "期望")
#: 长得像「期望」但不是取值的行（`期望来源` 是 C-138 那行文字说明）
EXP_ROW_EXCLUDE = ("期望来源",)


# ═════════════════════════ 一、TSV（Excel 剪贴板）═════════════════════════
def parse_tsv(text):
    """Excel 口径的 TSV 解析：`\\r\\n` / `\\r` / `\\n` 都算换行，制表符分列，**空串格保留**
    （中间的空格 = 「这一格不填」，不能被吃掉），只剔末尾的空行。

    返回 `[[格文本, …], …]`。与 `truth/model.py::paste_tsv` 开头那几行同一套规则。
    """
    rows = [ln.split("\t") for ln in
            str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    while rows and rows[-1] == [""]:
        rows.pop()
    return rows


def copy_tsv(model, indexes):
    """选区 → TSV（C-293/C-294 的复制半边）：按选区行列的**外接矩形**出，空格子留空串。

    与 `TruthModel.copy_tsv` 同行为（`test_copy_tsv_matches_model` 逐字比）。
    `indexes` 只当作「有 `.row()` / `.column()`（可能还有 `.isValid()`）的对象」，
    取值走 `model.data(ix)` 的默认 role（= DisplayRole），本模块不 import Qt。
    """
    cells = {}
    for ix in (indexes or ()):
        if ix is None:
            continue
        valid = getattr(ix, "isValid", None)
        if valid is not None and not valid():
            continue
        cells[(ix.row(), ix.column())] = model.data(ix) or ""
    if not cells:
        return ""
    rs = [r for (r, _c) in cells]
    cs = [c for (_r, c) in cells]
    lines = []
    for r in range(min(rs), max(rs) + 1):
        lines.append("\t".join(cells.get((r, c), "")
                               for c in range(min(cs), max(cs) + 1)))
    return "\n".join(lines)


#: 粘贴结果说明里最多点名几个「没认出写法」的格子（再多就 …，一行提示条塞不下）
BAD_NAMES_MAX = 6


@dataclass
class PasteTally:
    """跳过的格子按**原因**分桶（C-270 / C3-a 的护栏：跳过必有名字 + 原因，
    不能一律说成「只读」）。汇总那一句按这四桶逐条拼，见 `paste_report_text`。

    ro        只读格（auto_out 行 / 只读输入行 / iddq 自检拍列）
    no_col    落点右边没有测试列了，而且这个信号**加不出**新列
    mux_whole 自动生成列的 mux 数据值（C-110：只能整表一起改，逐格粘会把前一格覆盖掉）
    bad       写法没认出来的格子【名字】（`行标签×列名`，C-083 的粘贴面）
    """
    ro: int = 0
    no_col: int = 0
    mux_whole: int = 0
    bad: list = field(default_factory=list)
    #: R3-08：前三桶的**格名**（`行标签×列名`，与 `bad` 同一种写法）。
    #: 以前它们只有计数 —— 屏幕上写着「跳过 2 格（只读行/列）」，是哪两格没人说得出来。
    #: 计数仍由 `.n` 按四个数算（格名拿不到时照样报得出数，见 `paste_report_text`）。
    ro_cells: list = field(default_factory=list)
    no_col_cells: list = field(default_factory=list)
    mux_cells: list = field(default_factory=list)

    @property
    def n(self):
        """一共跳了几格。"""
        return int(self.ro) + int(self.no_col) + int(self.mux_whole) + len(self.bad)


@dataclass
class PastePlan:
    """一次粘贴的**规划**（不落格）。

    cells         [(行, 列, 该写的文本)] —— 能落的格子，**含「值没变」的那些**
                  （`len(cells)` 就是报告里那个「落了 N 格」，与 model 的 n_ok 同口径）
    new_cols      要先追加几条测试列（超出现有列数的部分，C-294「超列自动追加」）
    added_names   追加那几条列的【最终标号】**预测**（`plan_added_names`，与
                  `model.append_test_column` 的返回值相同；报告先按它报、落格后 model 用真值覆盖）
    kind          规划时的 `model.editable_kind()`（"" / "logic" / "mux"）——
                  「加不出新列」那句的原因按它分档
    rejected_rows 被拒的**源行**下标（超出真值表行数的那些）；非空 = 整次粘贴不做
    skipped       [(行, 列, 原因)] —— 逐格原因（给需要点名到格的调用方）
    tally         同一批跳过按原因分桶（汇总那一句按它拼）
    report_text   给用户看的那一句（`terms.TRUTH_PASTE_REPORT_FMT` / OVERFLOW）
    """
    cells: list = field(default_factory=list)
    new_cols: int = 0
    added_names: list = field(default_factory=list)
    kind: str = ""
    rejected_rows: list = field(default_factory=list)
    skipped: list = field(default_factory=list)
    tally: PasteTally = field(default_factory=PasteTally)
    report_text: str = ""

    @property
    def ok(self):
        """整次粘贴做不做得成（超行 = 整次拒绝，与 `model.paste_tsv` 的返回值同口径）。"""
        return not self.rejected_rows


def _cell_editable(model, r, c):
    """这一格能不能写 —— 与 `TruthModel._is_editable` 同判据，只经 Proto 的公开面。

    （`flags()` 是 Qt 的，本模块不 import Qt；`test_cell_editable_matches_model_flags`
    拿两张 mirror 的信号逐格对着 `model.flags()` 断言相等，漂了当场红。）
    """
    if not (0 <= c < model.columnCount()) or not (0 <= r < model.rowCount()):
        return False
    if not model.editable_kind():                       # C-134：不可建信号 → 全表只读
        return False
    kind = model.row_kind(r)
    if kind == contracts.TruthRowKind.AUTO:
        return False                                     # auto_out 行恒只读（C-077）
    if model.col_state(c) == contracts.TruthColState.DFT:
        return False                                     # iddq 自检拍列整列只读（C-129/C-130）
    if kind == contracts.TruthRowKind.EXP:
        return True
    return model.cell_state(r, c) == "editable"


def _new_col_editable(model, r, ref_col, new_is_dft=False):
    """**还没加出来的那条列**里第 r 行能不能写。

    输入行的可写性只跟【行】走（`edits.input_cell_editable` 只看输入行的角色），列只在
    「是不是 iddq 自检拍列」这一点上说话 —— 所以拿表上任一条非自检拍列当参照即可。

    表上一条列都没有时（清零过 → 零用例，然后从 Excel 整片粘回来）没有参照列可看：
    此时能加出列的只有 logic（`edits.add_col`；mux 的 `copy_cols` 没有源列可复制，
    压根加不出来），而 logic 的输入行**除 iddq 门行外全可编辑**（`truth/rows.py` 给
    每个分组 `editable=True`，只有门行是 False）—— 门行按行标签里的
    `inputs_table.ROLE_DFT_GATE` 认（那一行的标签就是 `vheader_display` 用这个常量拼的）。

    `new_is_dft` = 「新加出来那条列本身就是 iddq 自检拍列」（C-129/C-130 整列只读，
    **期望行也只读**）。mux 的加列 = 克隆**最后一列**（`edits.copy_cols` 不传 sel 时取
    `cols[-1]`），最后一列恰好是自检拍列时克隆出来的也是——C3-int 之前这里按「表上任一条
    非自检拍列」当参照，于是把这几格算成可写、model 落格时才发现整列只读，两边差几格。

    ⚠ 残留假设（logic）：`edits.add_col` 给新列的门值是 0，只有 iddq **透传值 = 1** 的信号
    会让它变成自检拍列（两张 mirror 上透传值都是 0，没有实例）。
    """
    if not model.editable_kind():
        return False
    kind = model.row_kind(r)
    if kind == contracts.TruthRowKind.AUTO:
        return False
    if new_is_dft:
        return False                      # C-129/C-130：自检拍列整列只读（期望行也是）
    if kind == contracts.TruthRowKind.EXP:
        return True
    if ref_col is not None:
        return model.cell_state(r, ref_col) == "editable"
    if model.columnCount() or model.editable_kind() != "logic":
        return False                      # 表上全是自检拍列 → 判不了，当只读（保守）
    return IT.ROLE_DFT_GATE not in str(model.row_label(r) or "")


def plan_added_names(model, n):
    """追加 n 条测试列会拿到的【最终标号】——`model.append_test_column(n)` 的**预测**。

    规划期报不出列名的话，那句「新增列 U0, U1」就只能事后补，而粘贴报告是一次性给的。
    预测规则 = `edits.new_col_name`（`truth_edit.uniq_col_name` 查全集）：
      · logic：`edits.add_col` 造正向列，名字就是 U<n>；
      · mux  ：`edits.copy_cols` 复制**最后一列**，负向标记跟着源列走（U<n>_NEG）；
        表上一条列都没有时 `copy_cols` 什么也造不出来 —— 预测同样给空，别报个假名字。
    `test_plan_added_names_matches_model` 拿 logic/mux 两条信号对着真 `append_test_column` 比。
    """
    n = int(n or 0)
    kind = model.editable_kind()
    if not kind or n <= 0:
        return []
    cols = list(model.cols())
    if kind == "mux" and not cols:
        return []                       # copy_cols 没有源列可复制 → 一条也加不出来
    raw = [str(c.get("name") or "") for c in cols]
    neg = bool(cols[-1].get("neg")) if (kind == "mux" and cols) else False
    out = []
    for _ in range(n):
        nm = TE.uniq_col_name(raw + out)
        out.append(nm)
    return [TE.final_col_name(nm, neg) for nm in out]


def paste_report_text(n_cells, added_names, skipped=0, kind=""):
    """粘贴结果那一句（C-294）—— **`truth/model.py` 与本模块共用这一份**（C3-int 去重）。

    `added_names` 给真实落下来的列名时文案就是最终态：`model._paste_land` 把
    `append_test_column` 的返回值喂进来，规划期那份预测名只用于 `plan.report_text`。

    `skipped` 两种都吃：
      · `int`        —— 只知道跳了几格（等价于全算「只读格」）；
      · `PasteTally` —— 四桶逐条点名（只读 / 右边没列（原因）/ mux 整表 / 写法没认出来 + 格名）。
    `kind` = `model.editable_kind()`，只用来分「加不出新列」那句的原因档（mux 是「清零后
    没有 case 可克隆」，logic 才是「这个信号加不出新列」）。
    """
    added = list(added_names or ())
    tally = skipped if isinstance(skipped, PasteTally) else PasteTally(ro=int(skipped or 0))
    tail = ""
    if tally.ro:
        tail += _named(terms.TRUTH_PASTE_SKIPPED_FMT, terms.TRUTH_PASTE_SKIPPED_PLAIN_FMT,
                       tally.ro_cells, tally.ro)
    if tally.no_col:
        why = (terms.TRUTH_PASTE_NO_NEW_COL_MUX if kind == "mux"
               else terms.TRUTH_PASTE_NO_NEW_COL)
        tail += _named(terms.TRUTH_PASTE_NO_COL_FMT, terms.TRUTH_PASTE_NO_COL_PLAIN_FMT,
                       tally.no_col_cells, tally.no_col, why=why)
    if tally.mux_whole:
        tail += _named(terms.TRUTH_PASTE_MUX_WHOLE_FMT, terms.TRUTH_PASTE_MUX_WHOLE_PLAIN_FMT,
                       tally.mux_cells, tally.mux_whole)
    if tally.bad:
        tail += terms.TRUTH_PASTE_BAD_FMT.format(n=len(tally.bad), names=_cells(tally.bad))
    return terms.TRUTH_PASTE_REPORT_FMT.format(
        n=int(n_cells),
        added=(terms.TRUTH_PASTE_ADDED_FMT.format(names=", ".join(added)) if added else ""),
        skipped=tail)


def _cell_name(model, names_after, r, c):
    """格名：`行标签×列名`（`terms.TRUTH_PASTE_BAD_CELL_FMT`）—— 四桶共用一种写法（R3-08）。"""
    return terms.TRUTH_PASTE_BAD_CELL_FMT.format(
        row=model.row_label(r), col=(names_after[c] if 0 <= c < len(names_after) else "?"))


def _cells(names):
    """格名列成一串（超出 `BAD_NAMES_MAX` 就省略号，别把整张表抄进一行）。"""
    shown = [str(x) for x in list(names)[:BAD_NAMES_MAX] if str(x).strip()]
    return "、".join(shown) + ("…" if len(list(names)) > BAD_NAMES_MAX else "")


def _named(fmt, plain_fmt, names, n, **kw):
    """点名版 / 不点名版 二选一（R3-08：点名在前、计数在后；拿不到格名时照实只报数）。"""
    text = _cells(names)
    if text:
        return fmt.format(cells=text, n=int(n), **kw)
    return plain_fmt.format(n=int(n), **kw)


def paste_plan(model, text, r0, c0):
    """TSV 粘贴的规划（C-294）：落在活动格 (r0, c0)；**超列追加列、超行整次拒绝**。

    超行为什么拒绝而不是追加：真值表的行是【输入信号】，凭空加一行等于凭空多一根输入，
    那是 Excel 表的事，不是这里能补的。跳过的格子**按四种原因分桶**（只读 / 右边没列 /
    自动生成列的 mux 数据值 / 写法没认出来），逐格给原因、汇总逐条点名，
    **整次粘贴照常落别的格**（C-083 的同一条原则：认不出来绝不静默吞成 0）。

    判据顺序与 `model._paste_cells` **必须一模一样**（`test_c294_paste_plan_equivalent_to_model_paste`
    逐格 + 逐字比）：没有列 → mux 整表数据值 → 只读 → 写法 → 落。

    返回 `PastePlan`。本函数**一格都不写**、一步撤销都不占。
    """
    kind = model.editable_kind()
    rows = parse_tsv(text)
    if not rows:
        return PastePlan(kind=kind, report_text=paste_report_text(0, [], 0, kind))
    r0, c0 = max(0, int(r0)), max(0, int(c0))
    n_rows = model.rowCount()
    if r0 + len(rows) > n_rows:
        # 拒掉的是【落不下的那些源行】—— 报给用户的是整次拒绝（与 model 同），
        # 但把具体哪几行落不下留在 plan 里，面板要点名时不用自己再算一遍。
        rejected = [i for i in range(len(rows)) if r0 + i >= n_rows]
        return PastePlan(kind=kind, rejected_rows=rejected,
                         report_text=terms.TRUTH_PASTE_OVERFLOW_ROWS_FMT.format(
                             rows=len(rows), max=n_rows))

    n_old = model.columnCount()
    need = c0 + max(len(x) for x in rows) - n_old
    new_cols = need if need > 0 else 0
    added = plan_added_names(model, new_cols)      # 真加得出来几条（mux 零列表时一条也加不出）
    n_have = n_old + len(added)                    # 追加之后一共有几列（落点判定按这个）
    ref_col = next((j for j in range(n_old)
                    if model.col_state(j) != contracts.TruthColState.DFT), None)
    # mux 的「加列」= 克隆最后一列（`edits.copy_cols` 不传 sel 时取 `cols[-1]`）：
    # 最后一列是 iddq 自检拍列时，克隆出来的整列只读（`edits.is_dft_pitch_col` 按门值判，
    # 门值是跟着 vals 一起复制走的）。不看这一条就会把新列算成可写、比 model 多落几格。
    new_is_dft = bool(kind == "mux" and n_old
                      and model.col_state(n_old - 1) == contracts.TruthColState.DFT)
    #: 追加之后的列名（写法没认出来的格子要报 `行标签×列名`，新列也得报得出名字）
    names_after = list(model.all_names()) + list(added)

    cells, skipped, tally = [], [], PasteTally()
    for dr, line in enumerate(rows):
        for dc, txt in enumerate(line):
            r, c = r0 + dr, c0 + dc
            if not (0 <= c < n_have):
                tally.no_col += 1
                tally.no_col_cells.append(_cell_name(model, names_after, r, c))
                skipped.append((r, c, terms.TRUTH_PASTE_SKIP_NO_COL))
                continue
            # 新追加的那几列是**手编列**（`edits.copy_cols` 置 user=True），
            # 按定义不会是「自动生成列的 mux 数据值」，所以只问表上已有的列。
            if c < n_old and model.is_mux_auto_data_cell(r, c):
                tally.mux_whole += 1
                tally.mux_cells.append(_cell_name(model, names_after, r, c))
                skipped.append((r, c, terms.TRUTH_PASTE_SKIP_MUX_WHOLE))
                continue
            can = (_cell_editable(model, r, c) if c < n_old
                   else _new_col_editable(model, r, ref_col, new_is_dft))
            if not can:
                tally.ro += 1
                tally.ro_cells.append(_cell_name(model, names_after, r, c))
                skipped.append((r, c, terms.TRUTH_PASTE_SKIP_READONLY))
                continue
            try:
                TE.parse_int(txt)                  # 空串 = 0 / 期望格的「清空」，都不算失败
            except ValueError:
                tally.bad.append(_cell_name(model, names_after, r, c))
                skipped.append((r, c, terms.TRUTH_PASTE_SKIP_PARSE_FMT.format(text=txt)))
                continue
            cells.append((r, c, txt))

    return PastePlan(cells=cells, new_cols=new_cols, added_names=list(added), kind=kind,
                     skipped=skipped, tally=tally,
                     report_text=paste_report_text(len(cells), added, tally, kind))


# ═════════════════════════ 二、导入期望（C-298）═════════════════════════
def _rows_from_csv(path):
    """CSV → 二维文本表。`utf-8-sig` 容忍 Excel 存出来的 BOM（导出那头写的就是它）。"""
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return [list(r) for r in csv.reader(f)]


def _rows_from_xlsx(path):
    """xlsx → 二维值表（**只读第一张表**）。openpyxl 惰性 import：没装它也不该拦住 CSV 那条路。"""
    import openpyxl                       # noqa: PLC0415  惰性：只有真读 xlsx 才需要
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _cell_text(v):
    """表格里一个格子 → 文本。xlsx 读出来的可能是 int/float（Excel 把 5 存成数字），
    那就直接当十进制值，别经 `parse_int` 的裸 hex 兜底把 11 读成 0x11。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v)


def _pick_exp_row(rows):
    """数据行里挑「期望」那一行。返回 (行, 用的哪一行的行名, 要不要提示)。

    找法（与导出的 CSV 行名对齐）：`期望(进.sv)` > `期望(bin)` > `期望` >（只有一行数据时）那一行。
    `期望来源` 长得像但是文字说明，明确排除掉。
    """
    data = rows[1:]
    for pre in EXP_ROW_PREFIXES:
        for r in data:
            lbl = _cell_text(r[0] if r else "").strip()
            if lbl in EXP_ROW_EXCLUDE:
                continue
            if lbl.startswith(pre):
                return r, lbl, None
    if len(data) == 1:
        lbl = _cell_text(data[0][0] if data[0] else "").strip()
        return data[0], lbl, terms.TRUTH_IMPORT_ONE_ROW_FMT.format(label=lbl)
    return None, "", terms.TRUTH_IMPORT_NO_EXP_ROW


def read_expectations(path):
    """从 CSV / xlsx 按【列名】读期望值（C-298 的读文件那半边）。

    返回 `(by_name, notes)`：

    * `by_name` = `{列名: int}` —— 直接喂 `model.apply_expectations(by_name)`。
      **给的是解析好的 int**，不是原文：`apply_expectations` 里是 `int(val) & mask`，
      而表里写的是 IC 工程师那 8 种写法（`0xA` / `4'b1010` / `16'hA` …），`int()` 一个都不认。
      文本 → 值的翻译是读文件这一侧的事（model 不碰文件、也不碰文本解析）。
    * `notes` = 逐条提示（哪一列的写法没认出来、列名重了取了哪一处、没找到期望行……）。
      **列名对不对得上本信号不在这里判** —— 那是 `model.apply_expectations` 返回的
      `missing`（C-298：报没对上的列名）。

    文件格式 = 导出的那张 CSV 原样改（`exports.signal_csv_text`）：第一行是列名
    （第一格是行名列的表头，如 `信号\\测试`），其后找「期望」行取值；整个文件只有一行
    数据时就按那一行取（designer 手写一张两行小表也能用）。
    """
    ext = os.path.splitext(str(path or ""))[1].lower()
    rows = _rows_from_xlsx(path) if ext in (".xlsx", ".xlsm") else _rows_from_csv(path)
    rows = [r for r in rows if any(_cell_text(v).strip() for v in (r or ()))]
    notes = []
    if not rows:
        return {}, [terms.TRUTH_IMPORT_EMPTY_FILE]
    header = [_cell_text(v).strip() for v in rows[0]]
    if len(header) < 2 or not any(header[1:]):
        return {}, [terms.TRUTH_IMPORT_NO_COLUMNS]

    row, _lbl, note = _pick_exp_row(rows)
    if note:
        notes.append(note)
    if row is None:
        return {}, notes

    by_name = {}
    for i, nm in enumerate(header):
        if i == 0 or not nm:
            continue
        if nm in by_name:
            notes.append(terms.TRUTH_IMPORT_DUP_COL_FMT.format(name=nm))
            continue
        txt = _cell_text(row[i] if i < len(row) else "").strip()
        if txt == "":
            continue                       # 这一列没填 = 不动它（不是「填 0」）
        try:
            by_name[nm] = TE.parse_int(txt)
        except ValueError:
            notes.append(terms.TRUTH_IMPORT_BAD_CELL_FMT.format(name=nm, text=txt))
    return by_name, notes


def import_report_text(n_applied, missing=(), notes=()):
    """导入期望的结果说明整句（C-298 / **I-20：先点名、计数在后**）。

    C3-int 之前这句有两份：`model.import_expectations` 拼的是「按列名回填了 N 列的期望；
    这些列名在本信号里没有：…」（计数在前，违反 I-20），面板自己又拼了一份名字在前的。
    现在只有这一份，两边都调：

        逐条原因（读文件时的 notes）；没对上的列名 + 共几个；按列名回填了 N 列的期望

    `missing` 为空、`notes` 为空时就只剩最后那半句。分号是 `terms` 之外的**结构符**，
    和 `dialogs` / 面板提示条的分段一致。
    """
    parts = [str(x) for x in (notes or ()) if str(x).strip()]
    miss = [str(x) for x in (missing or ())]
    if miss:
        parts.append(terms.TRUTH_IMPORT_MISSING_FMT.format(names="、".join(miss), n=len(miss)))
    head = ("；".join(parts) + "；") if parts else ""
    return terms.TRUTH_IMPORT_EXP_REPORT_FMT.format(missing=head, n=int(n_applied))


# ═════════════════════════ 三、单信号真值表 CSV（C-136…C-139）═════════════════════════
def csv_input_rows(an):
    """CSV 的输入行（第一列的信号/字段名）—— 与 v1 `SignalView.e_inputs` 的三个键同口径。

    行名用【带位宽的真名】，不带真值表冻结列那个 `　(角色 · 端口)` 后缀：CSV 是拿去和
    寄存器表 / .sv 对名字的交换件，多一截中文会让「按名字查」变难。
    行序与编辑器一致（logic 按 `an["groups"]`，mux 按 `expansion["used_vars"]`，iddq 门殿后）。
    """
    out = []
    an = an or {}
    kind = an.get("editable") or ""
    if kind == "logic":
        for g in (an.get("groups") or []):
            out.append({"key": g["key"], "label": g.get("label") or g.get("base") or g["key"],
                        "width": g.get("width") or 1})
    elif kind == "mux":
        exp = an.get("expansion") or {}
        bindings = exp.get("bindings") or {}
        for k in (exp.get("used_vars") or []):
            b = bindings.get(k)
            if b is None:
                continue
            out.append({"key": k, "label": IT.mux_label(b), "width": b.width})
    gate = an.get("dft_gate")
    if gate:
        out.append({"key": gate["key"], "label": gate["label"], "width": gate.get("width", 1)})
    return out


def _gate_value(gate, vec):
    """该向量把 iddq 门驱成了几（功能拍 = 透传值，DFT 自检拍 = 门被 force 的那个值）。
    与 `edits.cols_from_vectors` 同口径 —— 真值表上那一行显示什么，CSV 里就是什么。"""
    gv = gate["transp"]
    for (wl, wv, _ww) in (getattr(vec, "extra_forces", None) or []):
        if wl == gate["wire_lhs"]:
            gv = wv
    return gv


def _mux_columns_from_vectors(an, vecs):
    """build 的向量 → CSV 列模型（C-139 忠于产物）。"""
    exp = an.get("expansion") or {}
    out_w = int(an.get("out_width") or 1)
    cols, _rows = X.columns_from_mux_vectors(
        vecs, exp.get("bindings") or {}, exp.get("used_vars") or [],
        data_keys=exp.get("data_keys") or (), out_width=out_w)
    gate = an.get("dft_gate")
    if gate:
        # `columns_from_mux_vectors` 的 gate_value 是整表一个常量，对不上 DFT 自检拍
        # （那一拍门被 force 成非透传值）——按每条向量的 extra_forces 补，和编辑器一致。
        for col, v in zip(cols, vecs):
            col["vals"][gate["key"]] = _gate_value(gate, v)
    return cols


#: `cov` 认得的键 = `exports.render_sv` 的覆盖度形参（与 `sv_preview._cov_args()` 一一对应）
COV_KEYS = ("mode", "max_tests", "exhaustive", "sig_cov", "form_cov")


def _render_kwargs(cov):
    """`cov` → `exports.render_sv` 的关键字。None / {} → 用 render_sv 自己的默认。

    ⚠ 必须与算出 `an` 的那一档**同一个覆盖度**，否则 CSV 列集是一个档、编辑器是另一个档，
    看着像「导出漏了几条」。面板（C3-c）给的就是 `sv_preview._cov_args()` 那五个值
    （本信号：单点 > 逻辑类型 > 全局已由 `CoverageState.mode_for` 折平）。
    """
    if not cov:
        return {}
    bad = set(cov) - set(COV_KEYS)
    if bad:
        raise TypeError("cov 只认 %s，不认 %s" % (list(COV_KEYS), sorted(bad)))
    return {k: cov[k] for k in COV_KEYS if k in cov}


def signal_csv_columns(model, an, name, provider=None, edited=None, cov=None):
    """CSV 的列（每列一条测试）。

    mux 且给了 `provider`：走产物（C-139）—— `exports.render_sv(provider, only=[name],
    edited=edited)` 那次 build 里该信号实际渲染的向量。整信号负向、负向去重、T 编号重排、
    iddq 自检拍都在 build 里，编辑器那份列模型看不到。
    其余（logic / register / dft / 只读信号）：model 现在这张表 —— 所见即所得，
    也正是 v1 `SignalView.on_export_csv` 用的那份（`sv.cur_cols`）。

    mux 但**拿不到产物**（没传 provider / 该信号这次 build 里没产出块，比如被跳过或只记账）
    → 退回 model 的列。面板（C3-c）导出前该自己判一下并照实说一句，别让用户以为导出的
    CSV 与 .sv 对得上。
    """
    kw = _render_kwargs(cov)                   # cov 键写错是调用方的 bug，当场抛，不吞
    if (an or {}).get("editable") == "mux" and provider is not None:
        try:
            _text, build = X.render_sv(provider, only=[name], edited=edited, **kw)
        except Exception:                      # noqa: BLE001  导不出 .sv 不该连 CSV 也导不了
            build = None
        vecs = X.signal_build_vectors(build, name) if build else None
        if vecs:
            return _mux_columns_from_vectors(an, vecs)
    return list(model.cols())


def export_signal_csv(path, model, an, name, provider=None, edited=None, cov=None, cols=None):
    """导出【本信号】真值表 CSV（C-136），返回写出的文本。

    转置排版（第一列 = 信号/字段名，其后每列一条测试）与行序全在
    `exports.signal_csv_text`：输入行… → auto_out → 期望(进.sv) → **期望(bin)**（C-137）
    → **期望来源** → **负向?**（C-138）→ **force** → **RF_WRITE**（C-137）。

    `provider` / `edited` / `cov` 只有 mux 用得上（C-139，见 `signal_csv_columns`）：面板传
    `state.provider()`、`state.compute_edited()`、以及算 `an` 用的那一档覆盖度 —— 与
    .sv 预览 / 导出中心喂给 `exports.render_sv` 的是同几样东西，所以 CSV 与 .sv 必然同一份产物。

    `cols` 收**现成的列**（C3-int）：面板导出前要先问一次列才判得出「mux 这次有没有产物」
    （C-139 要照实说一句），传进来就不用再 `render_sv` 第二遍 —— 同一份列，写出来的
    CSV 与那次判断说的是同一件事。不传就自己问（`signal_csv_columns`）。
    """
    out_w = int((an or {}).get("out_width") or 1)
    if cols is None:
        cols = signal_csv_columns(model, an, name, provider=provider, edited=edited, cov=cov)
    cols = list(cols)
    rows = csv_input_rows(an)
    drive_fn = X.make_drive_fn(*X.drive_context(an))
    text = X.signal_csv_text(cols, rows, out_width=out_w, drive_fn=drive_fn)
    X.write_signal_csv(path, cols, rows, out_width=out_w, drive_fn=drive_fn,
                       name=name, text=text)
    return text
